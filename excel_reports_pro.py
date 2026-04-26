# -*- coding: utf-8 -*-
"""
excel_reports_pro.py
====================
Módulo de reportes Excel profesionales para:
  - Reporte de asistencia mensual por grupo
  - Reporte de horas y pagos mensual por docente
  - Reporte consolidado de todos los docentes (mes)

Uso desde app.py:
    from excel_reports_pro import generar_reporte_asistencia, generar_reporte_horas_docente, generar_reporte_consolidado_docentes

Cada función devuelve un objeto BytesIO listo para Flask send_file().
"""

import io
from datetime import timedelta, date
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.image import Image as XLImage


# ─────────────────────────────────────────────
#  PALETA DE COLORES INSTITUCIONAL
# ─────────────────────────────────────────────
C = {
    'azul_oscuro':   '1E3A5F',   # Cabeceras principales
    'azul_medio':    '2E6DA4',   # Cabeceras secundarias
    'azul_claro':    'D6E4F0',   # Filas alternas
    'verde':         '1A7A4A',   # Presente
    'verde_claro':   'C6EFCE',   # Fondo presente
    'rojo':          'C0392B',   # Ausente
    'rojo_claro':    'FFDEDE',   # Fondo ausente
    'amarillo':      'D4AC0D',   # Atraso
    'amarillo_claro':'FFF2CC',   # Fondo atraso
    'celeste':       '2980B9',   # Con permiso
    'celeste_claro': 'DDEEFF',   # Fondo con permiso
    'gris_header':   'F2F2F2',   # Fondo título meta
    'gris_texto':    '555555',   # Texto secundario
    'blanco':        'FFFFFF',
    'total_fondo':   '1E3A5F',   # Fila de totales
    'total_texto':   'FFFFFF',
    'naranja':       'E67E22',   # Alerta deuda
    'naranja_claro': 'FDEBD0',
}

MESES_ES = [
    '', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre'
]

ESTADOS_LABEL = {'P': 'Presente', 'A': 'Ausente', 'T': 'Atraso', 'N': 'Con permiso'}


# ─────────────────────────────────────────────
#  HELPERS DE ESTILO
# ─────────────────────────────────────────────

def _fill(hex_color):
    return PatternFill(fill_type='solid', fgColor=hex_color)

def _font(bold=False, size=11, color='000000', italic=False, name='Arial'):
    return Font(bold=bold, size=size, color=color, italic=italic, name=name)

def _border_thin():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

def _border_medium():
    s = Side(style='medium', color='888888')
    return Border(left=s, right=s, top=s, bottom=s)

def _align(h='center', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _set_col_width(ws, col_letter, width):
    ws.column_dimensions[col_letter].width = width

def _set_row_height(ws, row, height):
    ws.row_dimensions[row].height = height

def _style_cell(cell, fill=None, font=None, alignment=None, border=None):
    if fill:      cell.fill = fill
    if font:      cell.font = font
    if alignment: cell.alignment = alignment
    if border:    cell.border = border

def _merge_style(ws, cell_range, value, fill=None, font=None, alignment=None):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(':')[0]]
    cell.value = value
    _style_cell(cell, fill=fill, font=font, alignment=alignment)

def _bloque_info(ws, start_row, pairs, label_col=1, val_col=2):
    """Escribe pares (label, valor) en dos columnas con estilo limpio."""
    for i, (label, val) in enumerate(pairs):
        r = start_row + i
        lc = ws.cell(row=r, column=label_col, value=label)
        vc = ws.cell(row=r, column=val_col, value=val)
        lc.font = _font(bold=True, size=10, color=C['azul_oscuro'])
        lc.alignment = _align(h='right')
        vc.font = _font(size=10)
        vc.alignment = _align(h='left')
        lc.border = _border_thin()
        vc.border = _border_thin()
        if i % 2 == 0:
            lc.fill = _fill(C['gris_header'])
            vc.fill = _fill(C['gris_header'])


# ─────────────────────────────────────────────
#  PORTADA / ENCABEZADO PRINCIPAL
# ─────────────────────────────────────────────

def _encabezado_principal(ws, titulo, subtitulo, meta_pairs, ancho_total):
    """
    Dibuja la cabecera institucional de la hoja.
    Retorna el número de fila donde termina el bloque.
    """
    # Franja de color superior
    for col in range(1, ancho_total + 1):
        c = ws.cell(row=1, column=col)
        c.fill = _fill(C['azul_oscuro'])
    _set_row_height(ws, 1, 8)

    # Título principal
    end_col = get_column_letter(ancho_total)
    _merge_style(ws, f'A2:{end_col}4', titulo,
                 fill=_fill(C['azul_oscuro']),
                 font=_font(bold=True, size=18, color=C['blanco']),
                 alignment=_align(h='center', v='center'))
    _set_row_height(ws, 2, 10)
    _set_row_height(ws, 3, 32)
    _set_row_height(ws, 4, 10)

    # Subtítulo
    _merge_style(ws, f'A5:{end_col}6', subtitulo,
                 fill=_fill(C['azul_medio']),
                 font=_font(size=12, color=C['blanco'], italic=True),
                 alignment=_align(h='center', v='center'))
    _set_row_height(ws, 5, 8)
    _set_row_height(ws, 6, 22)
    _set_row_height(ws, 7, 8)

    # Bloque de metadatos en dos columnas
    mid = ancho_total // 2
    mid_col = get_column_letter(mid)
    end_col2 = get_column_letter(ancho_total)

    # Columna izquierda
    left_pairs = meta_pairs[:len(meta_pairs)//2 + len(meta_pairs)%2]
    right_pairs = meta_pairs[len(meta_pairs)//2 + len(meta_pairs)%2:]

    start_meta = 8
    for i, (label, val) in enumerate(left_pairs):
        r = start_meta + i
        lc = ws.cell(row=r, column=1, value=label)
        vc = ws.cell(row=r, column=2, value=val)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=mid)
        lc.font = _font(bold=True, size=10, color=C['azul_oscuro'], name='Arial')
        vc.font = _font(size=10, name='Arial')
        lc.alignment = _align(h='right', v='center')
        vc.alignment = _align(h='left', v='center')
        lc.border = _border_thin()
        vc.border = _border_thin()
        bg = C['gris_header'] if i % 2 == 0 else C['blanco']
        lc.fill = _fill(bg)
        vc.fill = _fill(bg)
        _set_row_height(ws, r, 18)

    for i, (label, val) in enumerate(right_pairs):
        r = start_meta + i
        lc = ws.cell(row=r, column=mid + 1, value=label)
        vc = ws.cell(row=r, column=mid + 2, value=val)
        ws.merge_cells(start_row=r, start_column=mid+2, end_row=r, end_column=ancho_total)
        lc.font = _font(bold=True, size=10, color=C['azul_oscuro'], name='Arial')
        vc.font = _font(size=10, name='Arial')
        lc.alignment = _align(h='right', v='center')
        vc.alignment = _align(h='left', v='center')
        lc.border = _border_thin()
        vc.border = _border_thin()
        bg = C['gris_header'] if i % 2 == 0 else C['blanco']
        lc.fill = _fill(bg)
        vc.fill = _fill(bg)

    end_meta = start_meta + max(len(left_pairs), len(right_pairs))
    # Franja separadora
    for col in range(1, ancho_total + 1):
        c = ws.cell(row=end_meta, column=col)
        c.fill = _fill(C['azul_claro'])
    _set_row_height(ws, end_meta, 6)

    return end_meta + 1


# ═══════════════════════════════════════════════════════════
#  1. REPORTE DE ASISTENCIA MENSUAL
# ═══════════════════════════════════════════════════════════

def generar_reporte_asistencia(grupo, estudiantes, mapa_asistencia, mes, anio):
    """
    Genera un Excel profesional de asistencia mensual.

    Parámetros:
        grupo              : objeto GrupoAsistencia (con .nombre, .categoria.nombre, .docente)
        estudiantes        : lista de objetos Cliente ordenados por nombre
        mapa_asistencia    : dict { cliente_id: { fecha: AsistenciaDia } }
        mes                : int (1-12)
        anio               : int

    Retorna: BytesIO
    """
    wb = Workbook()

    # ── Hoja 1: Asistencia diaria ──────────────────────────
    ws = wb.active
    ws.title = "Asistencia Mensual"
    ws.sheet_view.showGridLines = False

    primer_dia = date(anio, mes, 1)
    if mes == 12:
        ultimo_dia = date(anio + 1, 1, 1) - timedelta(days=1)
    else:
        ultimo_dia = date(anio, mes + 1, 1) - timedelta(days=1)

    dias = []
    d = primer_dia
    while d <= ultimo_dia:
        dias.append(d)
        d += timedelta(days=1)

    num_dias = len(dias)
    # Columnas: Nº | Apellidos | Nombre | 1..31 | P | A | T | N | % | Deuda
    ancho_total = 3 + num_dias + 6
    DIAS_NOMBRES = ['L','M','X','J','V','S','D']

    docente_nombre = (grupo.docente.nombre_completo
                      if hasattr(grupo, 'docente') and grupo.docente else 'Sin asignar')

    meta = [
        ('Categoría:',    grupo.categoria.nombre),
        ('Grupo / Curso:',grupo.nombre),
        ('Docente:',      docente_nombre),
        ('Mes:',          f'{MESES_ES[mes]} {anio}'),
        ('Período:',      f'{primer_dia.strftime("%d/%m/%Y")} — {ultimo_dia.strftime("%d/%m/%Y")}'),
        ('Estudiantes:',  str(len(estudiantes))),
    ]

    data_start = _encabezado_principal(
        ws,
        titulo=f'REPORTE DE ASISTENCIA  —  {MESES_ES[mes].upper()} {anio}',
        subtitulo=f'{grupo.categoria.nombre}  ›  {grupo.nombre}',
        meta_pairs=meta,
        ancho_total=ancho_total,
    )

    # ── Sub-encabezado de días de la semana ───────────────
    dias_col_start = 4  # columna donde empiezan los días

    # Fila de días del mes (número)
    r_dia_num = data_start
    ws.cell(row=r_dia_num, column=1, value='Nº').font       = _font(bold=True, size=9, color=C['blanco'], name='Arial')
    ws.cell(row=r_dia_num, column=2, value='Apellidos').font= _font(bold=True, size=9, color=C['blanco'], name='Arial')
    ws.cell(row=r_dia_num, column=3, value='Nombre').font   = _font(bold=True, size=9, color=C['blanco'], name='Arial')
    for col_idx in [1, 2, 3]:
        ws.merge_cells(start_row=r_dia_num, start_column=col_idx,
                       end_row=r_dia_num+1, end_column=col_idx)
        c = ws.cell(row=r_dia_num, column=col_idx)
        c.fill = _fill(C['azul_oscuro'])
        c.alignment = _align(v='center')
        c.border = _border_thin()

    for i, dia in enumerate(dias):
        col = dias_col_start + i
        c_num = ws.cell(row=r_dia_num, column=col, value=dia.day)
        c_num.font      = _font(bold=True, size=8, color=C['blanco'], name='Arial')
        c_num.fill      = _fill(C['azul_oscuro'])
        c_num.alignment = _align()
        c_num.border    = _border_thin()

        # Día de la semana
        c_dow = ws.cell(row=r_dia_num+1, column=col, value=DIAS_NOMBRES[dia.weekday()])
        c_dow.font      = _font(size=7, color=C['blanco'], name='Arial')
        c_dow.fill      = _fill(C['azul_medio'])
        c_dow.alignment = _align()
        c_dow.border    = _border_thin()

    # Columnas de resumen
    resumen_cols = ['P', 'A', 'T', 'N', '%', '⚠']
    resumen_labels = ['Pres.', 'Aus.', 'Atr.', 'Perm.', '% Asist.', 'Deuda']
    resumen_colors = [C['verde'], C['rojo'], C['amarillo'], C['celeste'], C['azul_medio'], C['naranja']]
    col_res_start = dias_col_start + num_dias

    for j, (label, color) in enumerate(zip(resumen_labels, resumen_colors)):
        col = col_res_start + j
        c = ws.cell(row=r_dia_num, column=col, value=label)
        c.font      = _font(bold=True, size=8, color=C['blanco'], name='Arial')
        c.fill      = _fill(color)
        c.alignment = _align(wrap=True)
        c.border    = _border_thin()
        ws.merge_cells(start_row=r_dia_num, start_column=col,
                       end_row=r_dia_num+1, end_column=col)

    _set_row_height(ws, r_dia_num, 20)
    _set_row_height(ws, r_dia_num+1, 14)

    # ── Filas de datos ────────────────────────────────────
    fills_estado = {
        'P': _fill(C['verde_claro']),
        'A': _fill(C['rojo_claro']),
        'T': _fill(C['amarillo_claro']),
        'N': _fill(C['celeste_claro']),
    }
    fonts_estado = {
        'P': _font(bold=True, size=8, color=C['verde'], name='Arial'),
        'A': _font(bold=True, size=8, color=C['rojo'], name='Arial'),
        'T': _font(bold=True, size=8, color=C['amarillo'], name='Arial'),
        'N': _font(bold=True, size=8, color=C['celeste'], name='Arial'),
    }
    ETIQUETAS_ESTADO = {'P': 'P', 'A': 'A', 'T': 'T', 'N': 'N'}

    data_row = r_dia_num + 2
    totales_col = {k: 0 for k in ['P', 'A', 'T', 'N']}

    for idx, est in enumerate(estudiantes):
        r = data_row + idx
        bg = C['azul_claro'] if idx % 2 == 0 else C['blanco']

        # Número
        cn = ws.cell(row=r, column=1, value=idx+1)
        cn.font = _font(size=9, name='Arial'); cn.alignment = _align(); cn.border = _border_thin()
        cn.fill = _fill(bg)

        # Apellidos y nombre
        apellido = getattr(est, 'apellido', '')
        nombre   = getattr(est, 'nombre', est.nombre_completo if hasattr(est, 'nombre_completo') else '')
        ca = ws.cell(row=r, column=2, value=apellido.upper())
        ca.font = _font(bold=True, size=9, name='Arial'); ca.alignment = _align(h='left'); ca.border = _border_thin(); ca.fill = _fill(bg)
        cn2 = ws.cell(row=r, column=3, value=nombre)
        cn2.font = _font(size=9, name='Arial'); cn2.alignment = _align(h='left'); cn2.border = _border_thin(); cn2.fill = _fill(bg)

        conteo = {'P': 0, 'A': 0, 'T': 0, 'N': 0}
        for i, dia in enumerate(dias):
            col = dias_col_start + i
            reg = mapa_asistencia.get(est.id, {}).get(dia)
            estado = reg.estado if reg else ''
            c = ws.cell(row=r, column=col, value=ETIQUETAS_ESTADO.get(estado, ''))
            if estado in fills_estado:
                c.fill  = fills_estado[estado]
                c.font  = fonts_estado[estado]
                conteo[estado] += 1
                totales_col[estado] += 1
            else:
                c.fill = _fill(bg)
                c.font = _font(size=8, color='BBBBBB', name='Arial')
            c.alignment = _align()
            c.border    = _border_thin()

        total = sum(conteo.values())
        pct   = round(conteo['P'] / total * 100, 1) if total else 0.0

        # Resumen
        vals_res = [conteo['P'], conteo['A'], conteo['T'], conteo['N'], pct]
        colors_res = [C['verde_claro'], C['rojo_claro'], C['amarillo_claro'], C['celeste_claro'], C['azul_claro']]
        fcolors_res = [C['verde'], C['rojo'], C['amarillo'], C['celeste'], C['azul_medio']]
        for j, (v, bg_r, fc) in enumerate(zip(vals_res, colors_res, fcolors_res)):
            col = col_res_start + j
            fmt = f'{v:.1f}%' if j == 4 else str(v)
            c = ws.cell(row=r, column=col, value=fmt)
            c.font = _font(bold=(j==4), size=9, color=fc, name='Arial')
            c.fill = _fill(bg_r)
            c.alignment = _align()
            c.border = _border_thin()

        # Columna deuda
        tiene_deuda = getattr(est, 'saldo_pendiente', 0) or 0
        col_deuda = col_res_start + 5
        if tiene_deuda and float(tiene_deuda) > 0:
            cd = ws.cell(row=r, column=col_deuda, value=f'${float(tiene_deuda):.2f}')
            cd.font = _font(bold=True, size=9, color=C['rojo'], name='Arial')
            cd.fill = _fill(C['naranja_claro'])
        else:
            cd = ws.cell(row=r, column=col_deuda, value='✓')
            cd.font = _font(size=9, color=C['verde'], name='Arial')
            cd.fill = _fill(C['verde_claro'])
        cd.alignment = _align()
        cd.border = _border_thin()

        _set_row_height(ws, r, 17)

    # ── Fila de totales ───────────────────────────────────
    r_total = data_row + len(estudiantes)
    ws.merge_cells(start_row=r_total, start_column=1, end_row=r_total, end_column=3)
    ct = ws.cell(row=r_total, column=1, value='TOTALES')
    ct.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
    ct.fill = _fill(C['azul_oscuro'])
    ct.alignment = _align()
    ct.border = _border_thin()

    for i in range(num_dias):
        c = ws.cell(row=r_total, column=dias_col_start+i)
        c.fill = _fill(C['azul_oscuro'])
        c.border = _border_thin()

    totales_vals = [totales_col['P'], totales_col['A'], totales_col['T'], totales_col['N']]
    total_general = sum(totales_vals)
    pct_general   = round(totales_col['P'] / total_general * 100, 1) if total_general else 0.0

    for j, v in enumerate(totales_vals + [pct_general]):
        col = col_res_start + j
        fmt = f'{v:.1f}%' if j == 4 else str(v)
        c = ws.cell(row=r_total, column=col, value=fmt)
        c.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
        c.fill = _fill(C['azul_oscuro'])
        c.alignment = _align()
        c.border = _border_thin()

    c_last = ws.cell(row=r_total, column=col_res_start+5)
    c_last.fill = _fill(C['azul_oscuro'])
    c_last.border = _border_thin()

    _set_row_height(ws, r_total, 22)

    # ── Leyenda de colores ────────────────────────────────
    r_leyenda = r_total + 2
    ws.cell(row=r_leyenda, column=1, value='LEYENDA:').font = _font(bold=True, size=9, name='Arial')
    leyenda_items = [
        ('P = Presente', C['verde_claro'], C['verde']),
        ('A = Ausente',  C['rojo_claro'],  C['rojo']),
        ('T = Atraso',   C['amarillo_claro'], C['amarillo']),
        ('N = Con permiso', C['celeste_claro'], C['celeste']),
        ('⚠ = Deuda pendiente', C['naranja_claro'], C['naranja']),
    ]
    for i, (text, bg, fc) in enumerate(leyenda_items):
        col = 2 + i * 2
        ws.merge_cells(start_row=r_leyenda, start_column=col, end_row=r_leyenda, end_column=col+1)
        c = ws.cell(row=r_leyenda, column=col, value=text)
        c.font = _font(bold=True, size=9, color=fc, name='Arial')
        c.fill = _fill(bg)
        c.alignment = _align()
        c.border = _border_thin()

    # ── Anchos de columna ────────────────────────────────
    _set_col_width(ws, 'A', 5)
    _set_col_width(ws, 'B', 18)
    _set_col_width(ws, 'C', 16)
    for i in range(num_dias):
        _set_col_width(ws, get_column_letter(dias_col_start+i), 3.2)
    col_widths_res = [5, 5, 5, 5, 8, 9]
    for j, w in enumerate(col_widths_res):
        _set_col_width(ws, get_column_letter(col_res_start+j), w)

    # ── Inmovilizar paneles ───────────────────────────────
    ws.freeze_panes = ws.cell(row=data_row, column=dias_col_start)

    # ── Hoja 2: Resumen estadístico ───────────────────────
    ws2 = wb.create_sheet("Resumen")
    ws2.sheet_view.showGridLines = False

    _encabezado_principal(
        ws2,
        titulo=f'RESUMEN DE ASISTENCIA  —  {MESES_ES[mes].upper()} {anio}',
        subtitulo=f'{grupo.categoria.nombre}  ›  {grupo.nombre}',
        meta_pairs=meta,
        ancho_total=8,
    )

    # Tabla de resumen por estudiante
    r_s = 16
    headers_s = ['Nº', 'Estudiante', 'Presente', 'Ausente', 'Atraso', 'Permiso', 'Total', '% Asistencia']
    h_colors  = [C['azul_oscuro']]*2 + [C['verde'], C['rojo'], C['amarillo'], C['celeste'], C['azul_medio'], C['azul_oscuro']]
    for j, (h, hc) in enumerate(zip(headers_s, h_colors)):
        c = ws2.cell(row=r_s, column=j+1, value=h)
        c.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
        c.fill = _fill(hc)
        c.alignment = _align(wrap=True)
        c.border = _border_thin()
    _set_row_height(ws2, r_s, 22)

    for idx, est in enumerate(estudiantes):
        r = r_s + 1 + idx
        bg = C['azul_claro'] if idx % 2 == 0 else C['blanco']
        conteo = {'P': 0, 'A': 0, 'T': 0, 'N': 0}
        for dia in dias:
            reg = mapa_asistencia.get(est.id, {}).get(dia)
            if reg and reg.estado in conteo:
                conteo[reg.estado] += 1
        total = sum(conteo.values())
        pct   = round(conteo['P'] / total * 100, 1) if total else 0.0
        nombre_completo = (est.nombre_completo if hasattr(est, 'nombre_completo')
                           else f'{getattr(est,"nombre","")} {getattr(est,"apellido","")}')
        vals = [idx+1, nombre_completo, conteo['P'], conteo['A'], conteo['T'], conteo['N'], total, f'{pct:.1f}%']
        for j, v in enumerate(vals):
            c = ws2.cell(row=r, column=j+1, value=v)
            c.font = _font(size=10, name='Arial')
            c.fill = _fill(bg)
            c.alignment = _align(h='center' if j != 1 else 'left')
            c.border = _border_thin()
        # Colorear % según umbral
        c_pct = ws2.cell(row=r, column=8)
        if pct >= 80:
            c_pct.font = _font(bold=True, size=10, color=C['verde'], name='Arial')
        elif pct >= 60:
            c_pct.font = _font(bold=True, size=10, color=C['amarillo'], name='Arial')
        else:
            c_pct.font = _font(bold=True, size=10, color=C['rojo'], name='Arial')
        _set_row_height(ws2, r, 18)

    # Totales
    r_tot2 = r_s + 1 + len(estudiantes)
    ws2.merge_cells(start_row=r_tot2, start_column=1, end_row=r_tot2, end_column=2)
    ct2 = ws2.cell(row=r_tot2, column=1, value='TOTALES')
    ct2.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
    ct2.fill = _fill(C['azul_oscuro'])
    ct2.alignment = _align()
    ct2.border = _border_thin()
    for j, v in enumerate([totales_col['P'], totales_col['A'], totales_col['T'], totales_col['N'],
                            total_general, f'{pct_general:.1f}%']):
        c = ws2.cell(row=r_tot2, column=3+j, value=v)
        c.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
        c.fill = _fill(C['azul_oscuro'])
        c.alignment = _align()
        c.border = _border_thin()

    # Anchos hoja 2
    for col_l, w in [('A',5),('B',28),('C',10),('D',10),('E',10),('F',10),('G',10),('H',14)]:
        _set_col_width(ws2, col_l, w)

    # ── Gráfico de pie en hoja 2 ──────────────────────────
    pie = PieChart()
    pie.title = "Distribución de Asistencia"
    pie.style = 10
    pie.width = 14
    pie.height = 10

    # Datos para el gráfico (fila de totales)
    data_ref   = Reference(ws2, min_col=3, min_row=r_tot2, max_col=6, max_row=r_tot2)
    pie.add_data(data_ref)
    pie.dataLabels = None

    # Colores del pie
    slice_colors = [C['verde'], C['rojo'], C['amarillo'], C['celeste']]
    for i, color in enumerate(slice_colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        if pie.series:
            pie.series[0].dPt.append(pt)

    ws2.add_chart(pie, f'A{r_tot2 + 3}')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ═══════════════════════════════════════════════════════════
#  2. REPORTE DE HORAS Y PAGOS — UN DOCENTE
# ═══════════════════════════════════════════════════════════

def generar_reporte_horas_docente(docente, registros, mes, anio):
    """
    Genera un Excel profesional de horas y pagos de un docente.

    Parámetros:
        docente   : objeto Docente
        registros : lista de RegistroHoras del mes
        mes       : int (1-12)
        anio      : int

    Retorna: BytesIO
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Horas y Pagos"
    ws.sheet_view.showGridLines = False

    total_horas = sum(r.horas_dictadas for r in registros)
    total_pago  = sum(r.pago_calculado for r in registros)

    meta = [
        ('Docente:',       docente.nombre_completo),
        ('Email:',         docente.email or '—'),
        ('Materia:',       docente.materia or '—'),
        ('Proyecto:',      docente.proyecto or '—'),
        ('Valor / hora:',  f'$ {docente.valor_hora:.2f}'),
        ('Mes:',           f'{MESES_ES[mes]} {anio}'),
        ('Total horas:',   f'{total_horas:.1f} h'),
        ('Total a pagar:', f'$ {total_pago:.2f}'),
    ]

    data_start = _encabezado_principal(
        ws,
        titulo=f'REPORTE DE HORAS Y PAGOS  —  {MESES_ES[mes].upper()} {anio}',
        subtitulo=f'Docente: {docente.nombre_completo}  ·  {docente.materia or ""}  ·  {docente.proyecto or ""}',
        meta_pairs=meta,
        ancho_total=6,
    )

    # ── Cabecera de tabla ─────────────────────────────────
    headers = ['Fecha', 'Día', 'Horas', 'Materia', 'Notas', 'Pago ($)']
    h_colors = [C['azul_oscuro'], C['azul_oscuro'], C['azul_medio'],
                C['azul_medio'],  C['azul_medio'],  C['verde']]

    r_header = data_start
    for j, (h, hc) in enumerate(zip(headers, h_colors)):
        c = ws.cell(row=r_header, column=j+1, value=h)
        c.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
        c.fill = _fill(hc)
        c.alignment = _align()
        c.border = _border_thin()
    _set_row_height(ws, r_header, 24)

    # ── Filas de datos ────────────────────────────────────
    DIAS_ES = ['Lunes','Martes','Miércoles','Jueves','Viernes','Sábado','Domingo']
    r_data = r_header + 1
    acumulado = 0.0

    for idx, reg in enumerate(sorted(registros, key=lambda x: x.fecha)):
        r = r_data + idx
        bg = C['azul_claro'] if idx % 2 == 0 else C['blanco']
        pago = reg.pago_calculado
        acumulado += pago

        vals = [
            reg.fecha.strftime('%d/%m/%Y'),
            DIAS_ES[reg.fecha.weekday()],
            reg.horas_dictadas,
            reg.materia or docente.materia or '—',
            reg.notas or '—',
            pago,
        ]
        fmts  = [None, None, '0.0', None, None, '"$"#,##0.00']
        aligns = ['center','center','center','left','left','center']

        for j, (v, fmt, al) in enumerate(zip(vals, fmts, aligns)):
            c = ws.cell(row=r, column=j+1, value=v)
            c.font = _font(size=10, name='Arial')
            c.fill = _fill(bg)
            c.alignment = _align(h=al, wrap=(j==4))
            c.border = _border_thin()
            if fmt:
                c.number_format = fmt

        # Resaltar fin de semana
        if reg.fecha.weekday() >= 5:
            for j in range(6):
                ws.cell(row=r, column=j+1).fill = _fill(C['amarillo_claro'])

        _set_row_height(ws, r, 18)

    # ── Fila de totales ───────────────────────────────────
    r_total = r_data + len(registros)
    ws.merge_cells(start_row=r_total, start_column=1, end_row=r_total, end_column=2)
    ct = ws.cell(row=r_total, column=1, value='TOTAL DEL MES')
    ct.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
    ct.fill = _fill(C['azul_oscuro'])
    ct.alignment = _align()
    ct.border = _border_thin()

    c_h = ws.cell(row=r_total, column=3, value=total_horas)
    c_h.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
    c_h.fill = _fill(C['azul_oscuro'])
    c_h.alignment = _align()
    c_h.number_format = '0.0'
    c_h.border = _border_thin()

    for j in [4, 5]:
        c = ws.cell(row=r_total, column=j)
        c.fill = _fill(C['azul_oscuro'])
        c.border = _border_thin()

    c_p = ws.cell(row=r_total, column=6, value=total_pago)
    c_p.font = _font(bold=True, size=12, color=C['blanco'], name='Arial')
    c_p.fill = _fill(C['verde'])
    c_p.alignment = _align()
    c_p.number_format = '"$"#,##0.00'
    c_p.border = _border_thin()
    _set_row_height(ws, r_total, 26)

    # ── Bloque de resumen financiero ──────────────────────
    r_fin = r_total + 2

    resumen_items = [
        ('Total horas trabajadas',  f'{total_horas:.1f} h',   C['azul_claro'],   C['azul_oscuro']),
        ('Valor por hora',          f'$ {docente.valor_hora:.2f}', C['azul_claro'], C['azul_oscuro']),
        ('TOTAL A PAGAR',           f'$ {total_pago:.2f}',    C['verde_claro'],  C['verde']),
    ]
    for i, (label, valor, bg, fc) in enumerate(resumen_items):
        r = r_fin + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        cl = ws.cell(row=r, column=1, value=label)
        cv = ws.cell(row=r, column=4, value=valor)
        cl.font = _font(bold=(i==2), size=11, color=C['azul_oscuro'], name='Arial')
        cv.font = _font(bold=True, size=12, color=fc, name='Arial')
        cl.fill = cv.fill = _fill(bg)
        cl.alignment = _align(h='right')
        cv.alignment = _align(h='center')
        cl.border = cv.border = _border_medium()
        _set_row_height(ws, r, 22)

    # ── Anchos ───────────────────────────────────────────
    for col_l, w in [('A',14),('B',12),('C',8),('D',22),('E',32),('F',14)]:
        _set_col_width(ws, col_l, w)

    # ── Hoja 2: Calendario visual del mes ────────────────
    ws2 = wb.create_sheet("Calendario")
    ws2.sheet_view.showGridLines = False
    _encabezado_principal(
        ws2,
        titulo=f'CALENDARIO DE CLASES  —  {MESES_ES[mes].upper()} {anio}',
        subtitulo=f'{docente.nombre_completo}',
        meta_pairs=meta[:4],
        ancho_total=7,
    )

    # Semanas del mes
    primer_dia = date(anio, mes, 1)
    if mes == 12:
        ultimo_dia = date(anio+1,1,1)-timedelta(days=1)
    else:
        ultimo_dia = date(anio,mes+1,1)-timedelta(days=1)

    dias_semana = ['Lunes','Martes','Miércoles','Jueves','Viernes','Sábado','Domingo']
    r_cal = 16
    for j, ds in enumerate(dias_semana):
        c = ws2.cell(row=r_cal, column=j+1, value=ds)
        c.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
        c.fill = _fill(C['azul_medio'] if j < 5 else C['gris_texto'])
        c.alignment = _align()
        c.border = _border_thin()
    _set_row_height(ws2, r_cal, 20)

    mapa_horas = {r.fecha: r for r in registros}
    cur = primer_dia - timedelta(days=primer_dia.weekday())
    r_week = r_cal + 1
    while cur <= ultimo_dia:
        for j in range(7):
            d = cur + timedelta(days=j)
            col = j + 1
            if d.month != mes:
                c = ws2.cell(row=r_week, column=col, value='')
                c.fill = _fill('F0F0F0')
                c.border = _border_thin()
            elif d in mapa_horas:
                reg = mapa_horas[d]
                text = f'{d.day}\n{reg.horas_dictadas}h\n${reg.pago_calculado:.2f}'
                c = ws2.cell(row=r_week, column=col, value=text)
                c.font = _font(bold=True, size=9, color=C['verde'], name='Arial')
                c.fill = _fill(C['verde_claro'])
                c.alignment = _align(wrap=True)
                c.border = _border_thin()
            else:
                c = ws2.cell(row=r_week, column=col, value=str(d.day))
                c.font = _font(size=10, color='999999', name='Arial')
                c.fill = _fill(C['blanco'])
                c.alignment = _align()
                c.border = _border_thin()
            _set_row_height(ws2, r_week, 40)
        cur += timedelta(days=7)
        r_week += 1

    for j in range(7):
        _set_col_width(ws2, get_column_letter(j+1), 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ═══════════════════════════════════════════════════════════
#  3. REPORTE CONSOLIDADO — TODOS LOS DOCENTES DEL MES
# ═══════════════════════════════════════════════════════════

def generar_reporte_consolidado_docentes(docentes_data, mes, anio):
    """
    Genera un Excel con el resumen de todos los docentes del mes.

    Parámetros:
        docentes_data: lista de dicts con claves:
            {
              'docente': objeto Docente,
              'registros': lista de RegistroHoras,
              'total_horas': float,
              'total_pago': float,
            }
        mes  : int
        anio : int

    Retorna: BytesIO
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    ws.sheet_view.showGridLines = False

    gran_total_horas = sum(d['total_horas'] for d in docentes_data)
    gran_total_pago  = sum(d['total_pago']  for d in docentes_data)

    meta = [
        ('Mes:',            f'{MESES_ES[mes]} {anio}'),
        ('Docentes:',       str(len(docentes_data))),
        ('Total horas:',    f'{gran_total_horas:.1f} h'),
        ('Total a pagar:',  f'$ {gran_total_pago:.2f}'),
    ]

    data_start = _encabezado_principal(
        ws,
        titulo=f'NÓMINA DE DOCENTES  —  {MESES_ES[mes].upper()} {anio}',
        subtitulo='Resumen consolidado de horas trabajadas y pagos',
        meta_pairs=meta,
        ancho_total=8,
    )

    # Cabecera
    headers = ['Nº', 'Docente', 'Materia', 'Proyecto', 'Val/h ($)', 'Horas', 'Días', 'Total ($)']
    h_colors = [C['azul_oscuro']]*4 + [C['azul_medio']]*3 + [C['verde']]
    r_h = data_start
    for j, (h, hc) in enumerate(zip(headers, h_colors)):
        c = ws.cell(row=r_h, column=j+1, value=h)
        c.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
        c.fill = _fill(hc)
        c.alignment = _align(wrap=True)
        c.border = _border_thin()
    _set_row_height(ws, r_h, 24)

    r_data = r_h + 1
    for idx, dd in enumerate(sorted(docentes_data, key=lambda x: x['docente'].apellido)):
        r = r_data + idx
        bg = C['azul_claro'] if idx % 2 == 0 else C['blanco']
        doc = dd['docente']
        vals = [
            idx + 1,
            doc.nombre_completo,
            doc.materia or '—',
            doc.proyecto or '—',
            doc.valor_hora,
            dd['total_horas'],
            len(dd['registros']),
            dd['total_pago'],
        ]
        fmts   = [None,None,None,None,'"$"#,##0.00','0.0',None,'"$"#,##0.00']
        aligns = ['center','left','left','left','center','center','center','center']

        for j, (v, fmt, al) in enumerate(zip(vals, fmts, aligns)):
            c = ws.cell(row=r, column=j+1, value=v)
            c.font = _font(size=10, name='Arial')
            c.fill = _fill(bg)
            c.alignment = _align(h=al)
            c.border = _border_thin()
            if fmt:
                c.number_format = fmt

        # Resaltar total
        ct = ws.cell(row=r, column=8)
        ct.font = _font(bold=True, size=11, color=C['verde'], name='Arial')
        ct.fill = _fill(C['verde_claro'])
        _set_row_height(ws, r, 20)

    # Fila gran total
    r_gt = r_data + len(docentes_data)
    ws.merge_cells(start_row=r_gt, start_column=1, end_row=r_gt, end_column=4)
    ct = ws.cell(row=r_gt, column=1, value='GRAN TOTAL')
    ct.font = _font(bold=True, size=12, color=C['blanco'], name='Arial')
    ct.fill = _fill(C['azul_oscuro'])
    ct.alignment = _align()
    ct.border = _border_thin()

    for j, (v, fmt) in enumerate([(None, None), (gran_total_horas,'0.0'),
                                   (len(docentes_data),None), (gran_total_pago,'"$"#,##0.00')]):
        col = 5 + j
        c = ws.cell(row=r_gt, column=col, value=v)
        c.font = _font(bold=True, size=12, color=C['blanco'], name='Arial')
        c.fill = _fill(C['azul_oscuro'] if j != 3 else C['verde'])
        c.alignment = _align()
        c.border = _border_thin()
        if fmt and v is not None:
            c.number_format = fmt
    _set_row_height(ws, r_gt, 28)

    for col_l, w in [('A',5),('B',26),('C',20),('D',20),('E',12),('F',10),('G',8),('H',14)]:
        _set_col_width(ws, col_l, w)

    # Hoja individual por docente
    for dd in docentes_data:
        doc = dd['docente']
        sheet_name = f"{doc.apellido[:12]}_{doc.nombre[:6]}"[:31]
        wsi = wb.create_sheet(sheet_name)
        wsi.sheet_view.showGridLines = False

        meta_i = [
            ('Docente:',  doc.nombre_completo),
            ('Materia:',  doc.materia or '—'),
            ('Proyecto:', doc.proyecto or '—'),
            ('Val/hora:', f'$ {doc.valor_hora:.2f}'),
            ('Mes:',      f'{MESES_ES[mes]} {anio}'),
            ('Horas:',    f'{dd["total_horas"]:.1f}'),
            ('Días trabajados:', str(len(dd['registros']))),
            ('TOTAL:',    f'$ {dd["total_pago"]:.2f}'),
        ]
        ds_i = _encabezado_principal(
            wsi,
            titulo=f'{doc.nombre_completo.upper()}  —  {MESES_ES[mes].upper()} {anio}',
            subtitulo=f'{doc.materia or ""}  ·  {doc.proyecto or ""}',
            meta_pairs=meta_i,
            ancho_total=5,
        )

        hdrs = ['Fecha', 'Horas', 'Materia', 'Notas', 'Pago ($)']
        hcs  = [C['azul_oscuro']]*2 + [C['azul_medio']]*2 + [C['verde']]
        for j, (h, hc) in enumerate(zip(hdrs, hcs)):
            c = wsi.cell(row=ds_i, column=j+1, value=h)
            c.font = _font(bold=True, size=10, color=C['blanco'], name='Arial')
            c.fill = _fill(hc)
            c.alignment = _align()
            c.border = _border_thin()
        _set_row_height(wsi, ds_i, 22)

        for idx2, reg in enumerate(sorted(dd['registros'], key=lambda x: x.fecha)):
            r = ds_i + 1 + idx2
            bg = C['azul_claro'] if idx2 % 2 == 0 else C['blanco']
            for j, v in enumerate([reg.fecha.strftime('%d/%m/%Y'), reg.horas_dictadas,
                                    reg.materia or doc.materia or '—', reg.notas or '—', reg.pago_calculado]):
                c = wsi.cell(row=r, column=j+1, value=v)
                c.font = _font(size=10, name='Arial')
                c.fill = _fill(bg)
                c.alignment = _align(h='center' if j in [0,1,4] else 'left', wrap=(j==3))
                c.border = _border_thin()
                if j in [1,4]:
                    c.number_format = '0.0' if j==1 else '"$"#,##0.00'
            _set_row_height(wsi, r, 18)

        r_tot_i = ds_i + 1 + len(dd['registros'])
        wsi.merge_cells(start_row=r_tot_i, start_column=1, end_row=r_tot_i, end_column=2)
        c = wsi.cell(row=r_tot_i, column=1, value='TOTAL')
        c.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
        c.fill = _fill(C['azul_oscuro'])
        c.alignment = _align()
        c.border = _border_thin()
        for j, (v, fmt) in enumerate([(dd['total_horas'],'0.0'),('',''),('',''),
                                       (dd['total_pago'],'"$"#,##0.00')]):
            col = 3 + j - 1
            cx = wsi.cell(row=r_tot_i, column=col+1, value=v if v != '' else None)
            cx.font = _font(bold=True, size=11, color=C['blanco'], name='Arial')
            cx.fill = _fill(C['verde'] if col+1 == 5 else C['azul_oscuro'])
            cx.alignment = _align()
            cx.border = _border_thin()
            if fmt:
                cx.number_format = fmt
        _set_row_height(wsi, r_tot_i, 24)

        for col_l, w in [('A',14),('B',8),('C',20),('D',30),('E',14)]:
            _set_col_width(wsi, col_l, w)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ─────────────────────────────────────────────
#  NOMBRES DE ARCHIVO SUGERIDOS
# ─────────────────────────────────────────────

def nombre_archivo_asistencia(grupo_nombre, mes, anio):
    return f'Asistencia_{grupo_nombre.replace(" ","_")}_{MESES_ES[mes]}_{anio}.xlsx'

def nombre_archivo_horas(docente_apellido, mes, anio):
    return f'Horas_{docente_apellido}_{MESES_ES[mes]}_{anio}.xlsx'

def nombre_archivo_consolidado(mes, anio):
    return f'Nomina_Docentes_{MESES_ES[mes]}_{anio}.xlsx'